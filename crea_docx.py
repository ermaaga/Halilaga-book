#!/usr/bin/env python3
"""Unisce i capitoli di romanzo/ in un unico file Word.

Uso:
    python3 crea_docx.py            # solo i capitoli già scritti
    python3 crea_docx.py --bozze    # include anche i capitoli vuoti (solo titolo)

Richiede python-docx (pip install python-docx).

Convenzioni dei file .md:
- la prima riga "# Titolo" è il titolo del capitolo;
- i paragrafi sono separati da una riga vuota;
- *corsivo* e **grassetto** sono supportati;
- una riga con solo *** è un cambio di scena;
- i commenti <!-- ... --> sono note di lavoro e non finiscono nel Word.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RADICE = Path(__file__).parent
CARTELLA = RADICE / "romanzo"
USCITA = RADICE / "Romanzo_Halil.docx"

TITOLO = "Il Mancino del Sultano"  # provvisorio: il titolo è ancora da decidere
SOTTOTITOLO = "La leggenda di Halil Halilaga"
AUTORE = "Ermal Halilaga"

# Pagina che introduce ogni parte, prima del capitolo indicato.
PARTI = {
    1: ("Parte prima", "Leshtej, 1787"),
    3: ("Parte seconda", "La guerra, 1787–1792"),
    7: ("Parte terza", "Il ritorno, 1792–1799"),
}

FONT = "Garamond"


def leggi_capitolo(path):
    testo = re.sub(r"<!--.*?-->", "", path.read_text(encoding="utf-8"), flags=re.S)
    righe = testo.strip().splitlines()
    titolo = righe[0].lstrip("#").strip() if righe and righe[0].startswith("#") else path.stem
    corpo = "\n".join(righe[1:]).strip()
    blocchi = [b.strip() for b in re.split(r"\n\s*\n", corpo) if b.strip()]
    return titolo, blocchi


def aggiungi_testo(par, testo):
    # **grassetto** e *corsivo*
    for pezzo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", testo):
        if not pezzo:
            continue
        if pezzo.startswith("**"):
            par.add_run(pezzo[2:-2]).bold = True
        elif pezzo.startswith("*"):
            par.add_run(pezzo[1:-1]).italic = True
        else:
            par.add_run(pezzo)


def numero_pagina(section):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tipo, testo in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = testo
        run._r.append(el)


def centrato(doc, testo, dimensione, prima=0, dopo=0, corsivo=False, grassetto=False):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Cm(0)
    par.paragraph_format.space_before = Pt(prima)
    par.paragraph_format.space_after = Pt(dopo)
    run = par.add_run(testo)
    run.font.size = Pt(dimensione)
    run.italic = corsivo
    run.bold = grassetto
    return par


def nuova_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main():
    bozze = "--bozze" in sys.argv
    doc = Document()

    # Formato libro (A5) e stile del testo
    sez = doc.sections[0]
    sez.page_width, sez.page_height = Cm(14.8), Cm(21)
    sez.top_margin = sez.bottom_margin = Cm(2)
    sez.left_margin = sez.right_margin = Cm(1.8)
    stile = doc.styles["Normal"]
    stile.font.name = FONT
    stile.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    stile.font.size = Pt(11.5)
    pf = stile.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(0.6)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.2

    # Frontespizio
    centrato(doc, TITOLO, 24, prima=150, dopo=12, grassetto=True)
    centrato(doc, SOTTOTITOLO, 13, dopo=60, corsivo=True)
    centrato(doc, AUTORE, 12)

    capitoli = sorted(CARTELLA.glob("*.md"))
    inclusi = 0
    for path in capitoli:
        titolo, blocchi = leggi_capitolo(path)
        if not blocchi and not bozze:
            continue
        n = int(path.stem[:2])

        if n in PARTI:
            nuova_pagina(doc)
            parte, sottotitolo = PARTI[n]
            centrato(doc, parte, 16, prima=150, dopo=8, grassetto=True)
            centrato(doc, sottotitolo, 12, corsivo=True)

        # Ogni capitolo in una nuova sezione: così il numero di pagina continua
        nuova = doc.add_section(WD_SECTION.NEW_PAGE)
        nuova.footer.is_linked_to_previous = True
        etichetta = "" if n == 0 or "epilogo" in path.stem else f"{n}"
        if etichetta:
            centrato(doc, etichetta, 12, prima=60, dopo=4)
        centrato(doc, titolo, 16, prima=0 if etichetta else 60, dopo=30, grassetto=True)

        primo = True
        for b in blocchi:
            if b == "***":
                centrato(doc, "*   *   *", 11, prima=10, dopo=10)
                primo = True
                continue
            par = doc.add_paragraph()
            if primo:  # niente rientro dopo il titolo o un cambio di scena
                par.paragraph_format.first_line_indent = Cm(0)
                primo = False
            aggiungi_testo(par, " ".join(b.splitlines()))
        if not blocchi:
            centrato(doc, "[da scrivere]", 11, corsivo=True)
        inclusi += 1

    numero_pagina(doc.sections[0])
    doc.save(USCITA)
    print(f"Creato {USCITA.name}: {inclusi} capitoli su {len(capitoli)}.")


if __name__ == "__main__":
    main()
